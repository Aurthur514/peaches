"""
Run an end-to-end smoke test using the internal heuristic LLM (no external providers).
"""
from io import BytesIO
import os
from task_3_glr_streamlit import extract_text_from_docx, extract_text_from_pdfs, find_placeholders, fill_docx_template
from task_3_internal_llm import propose_fields_from_template, map_fields_from_reports
from docx import Document


def run_internal_smoke(template_path, pdf_path, output_path):
    print(f"Reading template: {template_path}")
    with open(template_path, 'rb') as f:
        template_bytes = f.read()
    template_text, template_doc = extract_text_from_docx(BytesIO(template_bytes))

    print(f"Extracting text from PDF: {pdf_path}")
    with open(pdf_path, 'rb') as f:
        pdf_bytes = f.read()
    reports_text = extract_text_from_pdfs([BytesIO(pdf_bytes)])

    placeholders = find_placeholders(template_text)
    if placeholders:
        print(f"Found placeholders: {placeholders}")
    else:
        print("No placeholders found. Using internal heuristic to propose fields...")
        placeholders = propose_fields_from_template(template_text)
        print(f"Internal proposed fields: {placeholders}")

    print("Mapping fields using internal heuristics...")
    mapping, audit = map_fields_from_reports(placeholders, reports_text)
    print("Mapping:")
    for k, v in mapping.items():
        print(f"  {k}: {v}")
    print("Audit:")
    for k, v in audit.items():
        print(f"  {k}: rule={v.get('rule')}, snippet={v.get('snippet')}")

    print("Filling template...")
    filled = fill_docx_template(Document(BytesIO(template_bytes)), mapping)
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    filled.save(output_path)
    print(f"Saved filled document to: {output_path}")


if __name__ == '__main__':
    # Default example 1 paths
    tmpl = r"Task 3 - GLR Pipeline\Example 1 - USAA\Input\USAA 800 Claims GLR Template 4-24.docx"
    pdf = r"Task 3 - GLR Pipeline\Example 1 - USAA\Input\photo report.pdf"
    out = r"task_3_output\filled_example1_internal.docx"
    run_internal_smoke(tmpl, pdf, out)
