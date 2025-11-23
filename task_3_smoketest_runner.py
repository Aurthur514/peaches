"""
Smoke test runner for Task 3 GLR Pipeline.

Usage:
  python task_3_smoketest_runner.py --template "path/to/template.docx" --pdf "path/to/report.pdf" --output "path/to/output.docx"

This script imports functions from `task_3_glr_streamlit.py` and runs the flow
without a UI: extract text, ask LLM to map fields, fill the template, write output.
"""
import argparse
import os
from io import BytesIO

# Import functions from the Streamlit app
from task_3_glr_streamlit import (
    extract_text_from_docx,
    extract_text_from_pdfs,
    find_placeholders,
    ask_llm_for_fields,
    ask_llm_to_map,
    fill_docx_template,
)
from docx import Document


def run_smoke(template_path, pdf_path, output_path):
    print(f"Reading template: {template_path}")
    with open(template_path, 'rb') as f:
        template_bytes = f.read()
    template_text, template_doc = extract_text_from_docx(BytesIO(template_bytes))

    print(f"Extracting text from PDF: {pdf_path}")
    with open(pdf_path, 'rb') as f:
        pdf_bytes = f.read()
    reports_text = extract_text_from_pdfs([BytesIO(pdf_bytes)])

    placeholders = find_placeholders(template_text)
    if not placeholders:
        print("No placeholders found, asking LLM to propose fields...")
        placeholders = ask_llm_for_fields(template_text)
    else:
        print(f"Found placeholders: {placeholders}")

    print("Asking LLM to map fields to report values...")
    mapping = ask_llm_to_map(placeholders, reports_text)
    print("Mapping received:")
    print(mapping)

    print("Filling template...")
    filled = fill_docx_template(Document(BytesIO(template_bytes)), mapping)
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    filled.save(output_path)
    print(f"Saved filled document to: {output_path}")


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--template', required=True)
    parser.add_argument('--pdf', required=True)
    parser.add_argument('--output', required=True)
    args = parser.parse_args()
    run_smoke(args.template, args.pdf, args.output)
